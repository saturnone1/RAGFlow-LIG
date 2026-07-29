from __future__ import annotations

import base64
import hashlib
import io
import time
import zlib
from concurrent.futures import ThreadPoolExecutor
from zipfile import ZIP_STORED, ZipFile

import pytest
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import lig_runtime.document_password as document_password_module
from lig_runtime.document_key_rotation import audit_configs, keyring_fingerprint
from lig_runtime.document_password import (
    LEGACY_STORAGE_FIELD,
    PASSWORD_CLEAR_FIELD,
    PASSWORD_INPUT_FIELD,
    PASSWORD_STATUS_FIELD,
    PASSWORD_STORAGE_FIELD,
    DocumentPasswordError,
    encrypted_pdf_page_count,
    is_encrypted_pdf,
    normalize_document_password_for_binary,
    prepare_document_for_parsing,
    preserve_document_password_state,
    sanitize_parser_config,
    secure_parser_config,
    validate_document_access,
)


@pytest.fixture(autouse=True)
def password_key(monkeypatch):
    encoded = base64.b64encode(bytes(range(32))).decode("ascii")
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY", encoded)
    monkeypatch.setenv("LIG_PDF_PASSWORD_KEY", encoded)
    monkeypatch.delenv("LIG_DOCUMENT_PASSWORD_KEY_ID", raising=False)
    monkeypatch.delenv("LIG_DOCUMENT_PASSWORD_PREVIOUS_KEYS", raising=False)
    with document_password_module._decrypt_cache_lock:
        document_password_module._decrypt_cache.clear()
        document_password_module._decrypt_cache_size = 0
        document_password_module._decrypt_inflight.clear()


def encrypted_pdf(user_password: str, owner_password: str | None = None, algorithm: str = "AES-256") -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt(user_password, owner_password, algorithm=algorithm)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def plain_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def encrypted_text_pdf(user_password: str, owner_password: str | None = None) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (LIG encrypted PDF content) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.encrypt(user_password, owner_password)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def encrypted_office_zip(password: str, extension: str = "docx", label: str = "RAG-LIG") -> bytes:
    pytest.importorskip("msoffcrypto")
    from msoffcrypto.format.ooxml import OOXMLFile

    plain = io.BytesIO()
    if extension == "docx":
        pytest.importorskip("docx")
        from docx import Document

        document = Document()
        document.add_paragraph(f"{label} Word test")
        document.save(plain)
    elif extension == "xlsx":
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Name", "Value"])
        workbook.active.append([f"{label} Excel test", 42])
        workbook.active.append(["Second row", 84])
        workbook.save(plain)
    elif extension == "pptx":
        pytest.importorskip("pptx")
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"{label} PowerPoint test"
        presentation.save(plain)
    else:
        raise AssertionError(f"unsupported test extension: {extension}")
    encrypted = io.BytesIO()
    OOXMLFile(io.BytesIO(plain.getvalue())).encrypt(password, encrypted)
    return encrypted.getvalue()


def hwpx_document(password: str | None) -> bytes:
    section = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<hs:sec xmlns:hs="urn:hancom:section" xmlns:hp="urn:hancom:paragraph">'
        '<hp:p><hp:run><hp:t>암호화 HWPX 본문</hp:t></hp:run></hp:p>'
        '<hp:p><hp:run><hp:t>두 번째 문단</hp:t></hp:run></hp:p>'
        "</hs:sec>"
    ).encode()
    output = io.BytesIO()
    with ZipFile(output, "w", ZIP_STORED) as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        if password is None:
            archive.writestr("Contents/section0.xml", section)
            archive.writestr(
                "META-INF/manifest.xml",
                '<?xml version="1.0"?><manifest xmlns="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>',
            )
        else:
            salt = bytes(range(16))
            iv = bytes(range(16, 32))
            start_key = hashlib.sha256(password.encode()).digest()
            key = hashlib.pbkdf2_hmac("sha1", start_key, salt, 1024, dklen=32)
            compressor = zlib.compressobj(level=6, wbits=-15)
            compressed = compressor.compress(section) + compressor.flush()
            padded = compressed + bytes((-len(compressed)) % 16)
            encryptor = Cipher(algorithms.AES(key), modes.CBC(iv)).encryptor()
            ciphertext = encryptor.update(padded) + encryptor.finalize()
            archive.writestr("Contents/section0.xml", ciphertext)
            checksum = base64.b64encode(hashlib.sha256(section[:1024]).digest()).decode()
            manifest = f'''<?xml version="1.0"?>
                <odf:manifest xmlns:odf="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
                  <odf:file-entry full-path="Contents/section0.xml" media-type="application/xml" size="{len(section)}">
                    <odf:encryption-data checksum-type="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0#sha256-1k" checksum="{checksum}">
                      <odf:algorithm algorithm-name="http://www.w3.org/2001/04/xmlenc#aes256-cbc" initialisation-vector="{base64.b64encode(iv).decode()}"/>
                      <odf:key-derivation key-derivation-name="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0#pbkdf2" key-size="32" iteration-count="1024" salt="{base64.b64encode(salt).decode()}"/>
                      <odf:start-key-generation start-key-generation-name="http://www.w3.org/2000/09/xmldsig#sha256" key-size="32"/>
                    </odf:encryption-data>
                  </odf:file-entry>
                </odf:manifest>'''
            archive.writestr("META-INF/manifest.xml", manifest)
    return output.getvalue()


def secured(filename: str, password: str) -> dict:
    config, changed = secure_parser_config({PASSWORD_INPUT_FIELD: password}, filename)
    assert changed is True
    assert PASSWORD_INPUT_FIELD not in config
    assert password not in config[PASSWORD_STORAGE_FIELD]
    return config


def test_document_open_password_is_sealed_and_pdf_is_decrypted():
    encrypted = encrypted_pdf("열기-암호", "소유자-암호")
    config = secured("secure.pdf", "열기-암호")
    validate_document_access("secure.pdf", encrypted, config)
    assert encrypted_pdf_page_count("secure.pdf", encrypted, config) == 1

    name, plaintext, parser_config = prepare_document_for_parsing("secure.pdf", encrypted, config)
    assert name == "secure.pdf"
    assert PdfReader(io.BytesIO(plaintext)).is_encrypted is False
    assert PASSWORD_STORAGE_FIELD not in parser_config


def test_encrypted_pdf_detection_keeps_protected_files_out_of_repair_path():
    assert is_encrypted_pdf(encrypted_pdf("open-password")) is True
    assert is_encrypted_pdf(plain_pdf()) is False
    assert is_encrypted_pdf(b"not a pdf") is False


def test_plain_replacement_discards_submitted_or_stale_password_state():
    config = secured("plain.pdf", "no-longer-needed")

    assert validate_document_access("plain.pdf", plain_pdf(), config) is False
    normalized = normalize_document_password_for_binary("plain.pdf", plain_pdf(), config)

    assert PASSWORD_STORAGE_FIELD not in normalized
    assert PASSWORD_STATUS_FIELD not in normalized
    assert validate_document_access("secure.pdf", encrypted_pdf("needed"), secured("secure.pdf", "needed")) is True


def test_chunk_method_change_preserves_only_the_sealed_password_state():
    source = secured("secure.pdf", "keep-me")
    source[PASSWORD_STATUS_FIELD] = True
    target = {
        "layout_recognize": "PaddleOCR",
        PASSWORD_INPUT_FIELD: "must-not-survive",
        PASSWORD_STATUS_FIELD: True,
    }

    preserved = preserve_document_password_state(source, target)

    assert preserved["layout_recognize"] == "PaddleOCR"
    assert preserved[PASSWORD_STORAGE_FIELD] == source[PASSWORD_STORAGE_FIELD]
    assert PASSWORD_INPUT_FIELD not in preserved
    assert PASSWORD_STATUS_FIELD not in preserved


def test_pdf_decryption_rebuild_preserves_page_content_streams():
    encrypted = encrypted_text_pdf("open-password", "owner-password")

    _, plaintext, _ = prepare_document_for_parsing(
        "adobe-style.pdf", encrypted, secured("adobe-style.pdf", "open-password")
    )

    reader = PdfReader(io.BytesIO(plaintext))
    assert reader.is_encrypted is False
    assert b"LIG encrypted PDF content" in reader.pages[0].get_contents().get_data()
    assert "LIG encrypted PDF content" in reader.pages[0].extract_text()


@pytest.mark.parametrize("algorithm", ["RC4-128", "AES-128", "AES-256"])
def test_common_pdf_algorithms(algorithm):
    encrypted = encrypted_pdf("open-password", "owner-password", algorithm)
    name, plaintext, _ = prepare_document_for_parsing("secure.pdf", encrypted, secured("secure.pdf", "open-password"))
    assert name == "secure.pdf"
    assert PdfReader(io.BytesIO(plaintext)).is_encrypted is False


def test_pdf_wrong_missing_and_permissions_only_passwords():
    encrypted = encrypted_pdf("correct")
    with pytest.raises(DocumentPasswordError, match="incorrect"):
        validate_document_access("secure.pdf", encrypted, secured("secure.pdf", "wrong"))
    with pytest.raises(DocumentPasswordError, match="password protected"):
        validate_document_access("secure.pdf", encrypted, {})

    restricted = encrypted_pdf("", "owner-password")
    with pytest.raises(DocumentPasswordError, match="owner"):
        validate_document_access("restricted.pdf", restricted, secured("restricted.pdf", "not-owner"))
    validate_document_access("restricted.pdf", restricted, secured("restricted.pdf", "owner-password"))


@pytest.mark.parametrize("extension", ["docx", "xlsx", "pptx"])
def test_modern_office_password_decryption(extension):
    encrypted = encrypted_office_zip("office-password", extension)
    name, plaintext, parser_config = prepare_document_for_parsing(
        f"secure.{extension}", encrypted, secured(f"secure.{extension}", "office-password")
    )
    assert name == f"secure.{extension}"
    assert plaintext.startswith(b"PK")
    expected_part = {
        "docx": "word/document.xml",
        "xlsx": "xl/workbook.xml",
        "pptx": "ppt/presentation.xml",
    }[extension]
    with ZipFile(io.BytesIO(plaintext)) as archive:
        assert expected_part in archive.namelist()
    assert PASSWORD_STORAGE_FIELD not in parser_config


def test_office_wrong_or_missing_password():
    encrypted = encrypted_office_zip("office-password")
    with pytest.raises(DocumentPasswordError, match="password protected"):
        validate_document_access("secure.docx", encrypted, {})
    with pytest.raises(DocumentPasswordError, match="incorrect"):
        validate_document_access("secure.docx", encrypted, secured("secure.docx", "wrong"))


def test_encrypted_and_plain_hwpx_are_extracted_to_text():
    encrypted = hwpx_document("hwpx-password")
    name, plaintext, parser_config = prepare_document_for_parsing(
        "secure.hwpx", encrypted, secured("secure.hwpx", "hwpx-password")
    )
    assert name == "secure.hwpx.txt"
    assert plaintext.decode() == "암호화 HWPX 본문\n\n두 번째 문단"
    assert PASSWORD_STORAGE_FIELD not in parser_config

    plain_name, plain_text, _ = prepare_document_for_parsing("plain.hwpx", hwpx_document(None), {})
    assert plain_name == "plain.hwpx.txt"
    assert plain_text.decode() == "암호화 HWPX 본문\n\n두 번째 문단"

    stale_config = {PASSWORD_STORAGE_FIELD: "unreadable-stale-token"}
    normalized = normalize_document_password_for_binary("plain.hwpx", hwpx_document(None), stale_config)
    assert PASSWORD_STORAGE_FIELD not in normalized


def test_hwpx_wrong_or_missing_password():
    encrypted = hwpx_document("correct")
    with pytest.raises(DocumentPasswordError, match="password protected"):
        validate_document_access("secure.hwpx", encrypted, {})
    with pytest.raises(DocumentPasswordError, match="incorrect"):
        validate_document_access("secure.hwpx", encrypted, secured("secure.hwpx", "wrong"))


def test_status_redaction_clear_and_legacy_storage_migration():
    config = secured("secure.docx", "secret")
    sanitized = sanitize_parser_config(config)
    assert PASSWORD_STORAGE_FIELD not in sanitized
    assert sanitized[PASSWORD_STATUS_FIELD] is True

    cleared, changed = secure_parser_config({**config, PASSWORD_CLEAR_FIELD: True}, "secure.docx")
    assert changed is True
    assert cleared[PASSWORD_STORAGE_FIELD] is None

    legacy = {LEGACY_STORAGE_FIELD: config[PASSWORD_STORAGE_FIELD]}
    migrated, changed = secure_parser_config(legacy, "secure.pdf")
    assert changed is False
    assert LEGACY_STORAGE_FIELD not in migrated
    assert migrated[PASSWORD_STORAGE_FIELD] == config[PASSWORD_STORAGE_FIELD]

    duplicated, changed = secure_parser_config(
        {
            PASSWORD_STORAGE_FIELD: config[PASSWORD_STORAGE_FIELD],
            LEGACY_STORAGE_FIELD: "stale-legacy-token",
        },
        "secure.pdf",
    )
    assert changed is False
    assert LEGACY_STORAGE_FIELD not in duplicated
    assert duplicated[PASSWORD_STORAGE_FIELD] == config[PASSWORD_STORAGE_FIELD]


def test_password_rejected_for_unrelated_file_type_and_plain_office_is_unchanged():
    with pytest.raises(DocumentPasswordError, match="supported"):
        secure_parser_config({PASSWORD_INPUT_FIELD: "secret"}, "notes.txt")

    plain = io.BytesIO()
    with ZipFile(plain, "w") as archive:
        archive.writestr("word/document.xml", "plain")
    name, binary, config = prepare_document_for_parsing("plain.docx", plain.getvalue(), {})
    assert name == "plain.docx"
    assert binary == plain.getvalue()
    assert PASSWORD_STATUS_FIELD not in config


def test_public_password_status_is_never_stored_as_parser_state():
    config, changed = secure_parser_config({PASSWORD_STATUS_FIELD: True}, "plain.pdf")
    assert changed is False
    assert PASSWORD_STATUS_FIELD not in config


def test_resubmitting_the_same_password_keeps_the_active_token():
    config = secured("same.docx", "unchanged")
    updated, changed = secure_parser_config(
        {**config, PASSWORD_INPUT_FIELD: "unchanged"},
        "same.docx",
    )
    assert changed is False
    assert updated[PASSWORD_STORAGE_FIELD] == config[PASSWORD_STORAGE_FIELD]


def test_hwpx_resource_limits_reject_manifest_entry_count_and_pbkdf_work(monkeypatch):
    monkeypatch.setattr(document_password_module, "_MAX_HWPX_MANIFEST_SIZE", 32)
    with pytest.raises(DocumentPasswordError, match="manifest.xml is too large"):
        prepare_document_for_parsing("limited.hwpx", hwpx_document(None), {})

    monkeypatch.setattr(document_password_module, "_MAX_HWPX_MANIFEST_SIZE", 4 * 1024 * 1024)
    monkeypatch.setattr(document_password_module, "_MAX_HWPX_ARCHIVE_ENTRIES", 2)
    with pytest.raises(DocumentPasswordError, match="too many entries"):
        prepare_document_for_parsing("limited.hwpx", hwpx_document(None), {})

    monkeypatch.setattr(document_password_module, "_MAX_HWPX_ARCHIVE_ENTRIES", 4096)
    monkeypatch.setattr(document_password_module, "_MAX_HWPX_PBKDF_ITERATIONS", 100)
    with pytest.raises(DocumentPasswordError, match="parameters are outside"):
        prepare_document_for_parsing("limited.hwpx", hwpx_document("password"), secured("limited.hwpx", "password"))


def test_hwp_input_size_is_bounded_before_ole_parsing(monkeypatch):
    monkeypatch.setattr(document_password_module, "_MAX_HWP_FILE_SIZE", 4)
    with pytest.raises(DocumentPasswordError, match="too large"):
        prepare_document_for_parsing("oversized.hwp", b"12345", {})


def test_malformed_hwp_and_hwpx_are_rejected_cleanly():
    with pytest.raises(DocumentPasswordError, match="valid HWP 5"):
        prepare_document_for_parsing("damaged.hwp", b"not-an-ole-file", {})
    with pytest.raises(DocumentPasswordError, match="valid ZIP"):
        prepare_document_for_parsing("damaged.hwpx", b"not-a-zip-file", {})

    malformed_xml = io.BytesIO()
    with ZipFile(malformed_xml, "w", ZIP_STORED) as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", b"<broken>")
        archive.writestr(
            "META-INF/manifest.xml",
            '<?xml version="1.0"?><manifest xmlns="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>',
        )
    with pytest.raises(DocumentPasswordError, match="body XML is invalid"):
        prepare_document_for_parsing("damaged.hwpx", malformed_xml.getvalue(), {})


def test_hwpx_extracts_markdown_table_notes_and_embedded_object_inventory():
    section = b'''<?xml version="1.0" encoding="UTF-8"?>
      <hs:sec xmlns:hs="urn:hancom:section" xmlns:hp="urn:hancom:paragraph">
        <hp:p><hp:run><hp:t>Introduction</hp:t></hp:run></hp:p>
        <hp:tbl>
          <hp:tr><hp:tc><hp:p><hp:run><hp:t>Name</hp:t></hp:run></hp:p></hp:tc><hp:tc><hp:p><hp:run><hp:t>Value</hp:t></hp:run></hp:p></hp:tc></hp:tr>
          <hp:tr><hp:tc><hp:p><hp:run><hp:t>Alpha</hp:t></hp:run></hp:p></hp:tc><hp:tc><hp:p><hp:run><hp:t>42</hp:t></hp:run></hp:p></hp:tc></hp:tr>
        </hp:tbl>
        <hp:p><hp:footNote><hp:run><hp:t>Reference note</hp:t></hp:run></hp:footNote></hp:p>
      </hs:sec>'''
    output = io.BytesIO()
    with ZipFile(output, "w", ZIP_STORED) as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", section)
        archive.writestr(
            "META-INF/manifest.xml",
            '<?xml version="1.0"?><manifest xmlns="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>',
        )
        archive.writestr("BinData/image.png", b"image")

    _, plaintext, _ = prepare_document_for_parsing("structured.hwpx", output.getvalue(), {})
    text = plaintext.decode()
    assert "| Name | Value |" in text
    assert "| Alpha | 42 |" in text
    assert "[Footnote] Reference note" in text
    assert "[Embedded object: BinData/image.png]" in text


def test_plain_office_is_not_cached_and_hwpx_preparation_is_single_flight(monkeypatch):
    with document_password_module._decrypt_cache_lock:
        document_password_module._decrypt_cache.clear()
        document_password_module._decrypt_cache_size = 0
        document_password_module._decrypt_inflight.clear()

    plain_office = io.BytesIO()
    with ZipFile(plain_office, "w") as archive:
        archive.writestr("word/document.xml", "plain")
    prepare_document_for_parsing("plain.docx", plain_office.getvalue(), {})
    assert not document_password_module._decrypt_cache

    binary = hwpx_document(None)
    original = document_password_module._prepare_hwpx
    calls = 0

    def counted_prepare(filename, payload, config):
        nonlocal calls
        calls += 1
        time.sleep(0.05)
        return original(filename, payload, config)

    monkeypatch.setattr(document_password_module, "_prepare_hwpx", counted_prepare)
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(lambda _: prepare_document_for_parsing("parallel.hwpx", binary, {}), range(8)))
    assert calls == 1
    assert len({result[1] for result in results}) == 1


def test_decryption_cache_entries_expire(monkeypatch):
    now = 100.0
    monkeypatch.setattr(document_password_module, "monotonic", lambda: now)
    binary = hwpx_document(None)
    prepare_document_for_parsing("ttl.hwpx", binary, {})
    key = document_password_module._cache_key(binary, None, ".hwpx")
    assert document_password_module._cached_plaintext(key) is not None
    now += document_password_module._DECRYPT_CACHE_TTL_SECONDS + 1
    assert document_password_module._cached_plaintext(key) is None


def test_key_rotation_rewraps_with_active_key_and_keeps_document_readable(monkeypatch):
    old_key = base64.b64encode(bytes(range(32))).decode()
    new_key = base64.b64encode(bytes(reversed(range(32)))).decode()
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY_ID", "old")
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY", old_key)
    config = secured("rotated.docx", "rotation-password")
    assert config[PASSWORD_STORAGE_FIELD].startswith("v3:old:")

    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY_ID", "new")
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY", new_key)
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_PREVIOUS_KEYS", f"old:{old_key}")
    rewrapped, changed = document_password_module.rewrap_parser_config(config)
    assert changed is True
    assert rewrapped[PASSWORD_STORAGE_FIELD].startswith("v3:new:")
    assert document_password_module._unseal(rewrapped[PASSWORD_STORAGE_FIELD]) == "rotation-password"

    monkeypatch.delenv("LIG_DOCUMENT_PASSWORD_PREVIOUS_KEYS")
    assert document_password_module._unseal(rewrapped[PASSWORD_STORAGE_FIELD]) == "rotation-password"


def test_key_rotation_requires_the_referenced_previous_key(monkeypatch):
    old_key = base64.b64encode(bytes(range(32))).decode()
    new_key = base64.b64encode(bytes(reversed(range(32)))).decode()
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY_ID", "old")
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY", old_key)
    config = secured("rotated.docx", "rotation-password")

    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY_ID", "new")
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY", new_key)
    with pytest.raises(DocumentPasswordError, match="is not configured"):
        document_password_module.rewrap_parser_config(config)


def test_key_rotation_migrates_v2_tokens_and_audits_bad_records(monkeypatch):
    old_key_bytes = bytes(range(32))
    old_key = base64.b64encode(old_key_bytes).decode()
    new_key = base64.b64encode(bytes(reversed(range(32)))).decode()
    nonce = bytes(range(12))
    ciphertext = AESGCM(old_key_bytes).encrypt(
        nonce,
        b"legacy-v2-password",
        document_password_module._AAD,
    )
    legacy_config = {
        PASSWORD_STORAGE_FIELD: "v2:" + base64.b64encode(nonce + ciphertext).decode()
    }

    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY_ID", "new")
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_KEY", new_key)
    monkeypatch.setenv("LIG_DOCUMENT_PASSWORD_PREVIOUS_KEYS", f"old:{old_key}")
    updates, errors = audit_configs(
        [("legacy", legacy_config), ("bad", {PASSWORD_STORAGE_FIELD: "v3:missing:not-base64"})]
    )

    assert errors and errors[0].startswith("bad:")
    assert updates[0][0] == "legacy"
    assert updates[0][1][PASSWORD_STORAGE_FIELD].startswith("v3:new:")
    assert document_password_module._unseal(updates[0][1][PASSWORD_STORAGE_FIELD]) == "legacy-v2-password"
    assert keyring_fingerprint()["configured_key_ids"] == ["new", "old"]
